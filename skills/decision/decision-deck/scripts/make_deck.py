#!/usr/bin/env python3
"""Build a 16:9 deck from deck.yaml. Usage: make_deck.py deck.yaml out/deck.pptx"""
import os, shutil, subprocess, sys
import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

INK = RGBColor(0x15, 0x20, 0x2B)
MUTED = RGBColor(0x5A, 0x68, 0x76)
ACCENT = RGBColor(0x0F, 0x7F, 0x86)
W, H = Inches(13.333), Inches(7.5)
M = Inches(0.6)


def text(slide, x, y, w, h, s, size, color=INK, bold=False):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text, p.font.size, p.font.bold, p.font.color.rgb = s, Pt(size), bold, color
    return tf


def bullets(slide, x, y, w, h, items):
    tf = slide.shapes.add_textbox(x, y, w, h).text_frame
    tf.word_wrap = True
    for i, b in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text, p.font.size, p.font.color.rgb = "• " + str(b), Pt(18), INK
        p.space_after = Pt(10)


def table(slide, x, y, w, rows):
    n, m = len(rows), len(rows[0])
    t = slide.shapes.add_table(n, m, x, y, w, Inches(0.45) * n).table
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = t.cell(i, j)
            c.text = str(v)
            for p in c.text_frame.paragraphs:
                p.font.size, p.font.bold = Pt(14), i == 0


def figure(slide, path, x, y, w, h):
    from PIL import Image
    iw, ih = Image.open(path).size
    s = min(w / iw, h / ih)
    slide.shapes.add_picture(path, x + int((w - iw * s) / 2), y, int(iw * s), int(ih * s))


def main(spec_path, out_path):
    spec = yaml.safe_load(open(spec_path))
    root = os.getcwd()
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    text(s, M, Inches(2.4), W - 2 * M, Inches(1.6), spec["title"], 36, bold=True)
    text(s, M, Inches(4.2), W - 2 * M, Inches(0.8), spec.get("subtitle", ""), 18, MUTED)

    for sl in spec["slides"]:
        s = prs.slides.add_slide(blank)
        text(s, M, Inches(0.4), W - 2 * M, Inches(1.1), sl["title"], 26, bold=True)
        s.shapes.add_connector(1, M, Inches(1.5), W - M, Inches(1.5)).line.color.rgb = ACCENT
        top, body_h = Inches(1.8), H - Inches(2.4)
        has_side = "figure" in sl or "table" in sl
        bw = Inches(5.2) if has_side else W - 2 * M
        if sl.get("bullets"):
            bullets(s, M, top, bw, body_h, sl["bullets"])
        sx = M + bw + Inches(0.3) if sl.get("bullets") else M
        sw = W - sx - M
        if "figure" in sl:
            figure(s, os.path.join(root, sl["figure"]), sx, top, sw, body_h)
        elif "table" in sl:
            table(s, sx, top, sw, sl["table"])
        if sl.get("note"):
            s.notes_slide.notes_text_frame.text = sl["note"]

    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    prs.save(out_path)
    print("saved", out_path, "slides:", len(prs.slides))
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if office:
        subprocess.run([office, "--headless", "--convert-to", "pdf", "--outdir",
                        os.path.dirname(out_path) or ".", out_path],
                       check=False, capture_output=True, timeout=180)
        print("pdf", os.path.splitext(out_path)[0] + ".pdf")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
